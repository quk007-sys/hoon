"""견적서 양식(참조 xlsx)을 기반으로 가스안전 AI 어드바이저 2.0 고도화 견적서 생성.
usage: python build_quote.py <template.xlsx> <out.xlsx>
"""
import sys, copy, datetime
from openpyxl import load_workbook

SRC, OUT = sys.argv[1], sys.argv[2]
TARGET = 1_000_000_000  # 목표 견적가(VAT 별도)

wb = load_workbook(SRC)
ws = wb.active

# ---- 양식 스타일 원형 수집 (행 단위, A~K) ----
COLS = 'ABCDEFGHIJK'
def grab(row):
    return {c: copy.copy(ws[f'{c}{row}']._style) for c in COLS}
P = {
    'sec': grab(37),        # 섹션 제목
    'sum_hdr': grab(15), 'sum_hw': grab(16), 'sum_row': grab(17), 'sum_tot': grab(21),
    'hw_hdr': grab(25), 'hw_row': grab(26), 'hw_tot': grab(30),
    'dv_hdr': grab(38), 'dv_first': grab(39), 'dv_row': grab(40), 'dv_tot': grab(44),
    'note_t': grab(83), 'note': grab(84), 'blank': grab(14),
}
HW_ROWS = [[ws.cell(row=r, column=ci).value for ci in range(1, 12)] for r in range(26, 30)]

# ---- 13행 이하 초기화 ----
for rng in list(ws.merged_cells.ranges):
    if rng.min_row >= 13:
        ws.unmerge_cells(str(rng))
for r in range(13, ws.max_row + 1):
    ws.row_dimensions[r].height = None
    ws.row_dimensions[r].hidden = False
    for ci in range(1, 31):
        cell = ws.cell(row=r, column=ci)
        cell.value = None
        cell._style = copy.copy(P['blank']['A'])

ws['K4'] = datetime.datetime(2026, 10, 8)

def style_row(r, proto, height=None):
    for c in COLS:
        ws[f'{c}{r}']._style = copy.copy(P[proto][c])
    if height:
        ws.row_dimensions[r].height = height

def gap(r, h=3.0):
    ws.row_dimensions[r].height = h

# ---- 데이터 ----
RATE = {  # 양식 기재 월평균 임금(원)
    'IT PM': 10086804, 'IT 아키텍트': 11103230, 'IT 기획자': 11853218,
    '데이터분석가': 8499309, '응용SW 개발자': 7754124, 'UI/UX기획/개발자': 6901660,
}
DEV = [
    ('과제 1. 가스안전 지식DB 고도화', '가스안전 지식DB 고도화', [
        ('IT PM', '사업관리 및 데이터 컨설팅 총괄', 1.0),
        ('IT 아키텍트', '데이터 표준·메타데이터 스키마·RAG 파이프라인 설계', 1.5),
        ('데이터분석가', '비정형 자료 전수조사·정제·분류·최신화 체계 수립', 4.0),
        ('응용SW 개발자', '수집·변환 파이프라인, 자동 업로드·버전관리, 관리자 포털 개선', 4.0),
    ]),
    ('과제 2. 스마트온2.0 연계 모바일', '스마트온2.0 연계 모바일', [
        ('IT PM', '사업관리 및 현업(검사지원처) 협의', 0.5),
        ('IT 기획자', '모바일 연계·계정/인증 보안 아키텍처 설계', 1.0),
        ('응용SW 개발자', '스마트온2.0 연계 API·SSO 연동·음성입력(STT) 개발', 3.5),
        ('UI/UX기획/개발자', '현장검사원 특화 모바일 UI/UX 설계·구현', 1.5),
    ]),
    ('과제 3. 사이버지사 VOC 챗봇(대국민)', '사이버지사 VOC 챗봇(대국민)', [
        ('IT PM', '사업관리 및 대외 협의(홍보소통실)', 0.5),
        ('IT 기획자', '대국민 상담체계(자동응답·근거답변·부서이관) 기획', 1.5),
        ('IT 아키텍트', '외부망 보안 아키텍처·내부자료 분리(필터링) 설계', 1.0),
        ('데이터분석가', '공개자료(법령·KGS Code) 선별·정제, 대국민 지식DB 구축', 1.0),
        ('응용SW 개발자', 'AI상담 챗봇·질의 라우팅·전문질의 이관/VOC 연계·가드레일 개발', 4.0),
        ('UI/UX기획/개발자', '대국민 챗봇 UI/UX(웹접근성) 설계·화면 구현', 2.0),
    ]),
    ('과제 4. AI 신뢰성(정확도) 고도화', 'AI 신뢰성(정확도) 고도화', [
        ('IT PM', '사업관리 및 품질관리', 1.0),
        ('IT 기획자', '관리지표(KPI)·평가체계·정확도 개선 로드맵 수립', 1.0),
        ('데이터분석가', '가스안전 전문 평가셋 구축·정량 평가·오답 분석', 3.0),
        ('응용SW 개발자', 'RAG 검색·청킹·Reranking 고도화, 근거제시·할루시네이션 가드레일, 성능평가 자동화', 5.0),
    ]),
]
SW = [  # 용도, 제품명, 상세내역, 납품단가(원), 수량
    ('문서보안(DRM)', 'DRM 보안소프트웨어', ' - 파일 복·암호화 솔루션\n - 지식DB 수집문서 복호화 연동', 14500000, 1),
    ('통합인증(SSO)', 'SSO 인증서버\n소프트웨어', ' - 기존 어드바이저 통합인증(로그인)\n - 스마트온2.0 계정·인증 연계', 20000000, 1),
    ('음성인식(STT)', 'STT 엔진\n(온프레미스)', ' - 현장 음성질의 입력\n - 가스안전 전문용어 사전 적용', 30000000, 1),
]

# ---- 레이아웃 계산 ----
r = 13
ws[f'A{r}'] = '총 견적내역'; ws[f'A{r}']._style = copy.copy(P['sec']['A'])
SUM_HDR = 15; gap(14, 7.5)
SUM_ROWS = list(range(16, 16 + 1 + 1 + len(DEV) + 1))  # HW, SW, 과제1~4, 단수조정
SUM_TOT = SUM_ROWS[-1] + 1
r = SUM_TOT + 1; gap(r); gap(r + 1); r += 2

# 1. HW
HW_SEC = r; HW_HDR = r + 1; HW_FIRST = r + 2; HW_TOT = HW_FIRST + len(HW_ROWS)
r = HW_TOT + 1; gap(r); gap(r + 1); r += 2
# 2. SW
SW_SEC = r; SW_HDR = r + 1; SW_FIRST = r + 2; SW_TOT = SW_FIRST + len(SW)
r = SW_TOT + 1; gap(r); gap(r + 1); r += 2
# 3~6. 개발비
dev_layout = []
for title, label, roles in DEV:
    sec, hdr, first = r, r + 1, r + 2
    tot = first + len(roles)
    dev_layout.append((sec, hdr, first, tot))
    r = tot + 1; gap(r); gap(r + 1); r += 2
NOTE = r + 1

# ---- 총 견적내역 ----
style_row(SUM_HDR, 'sum_hdr', 18.75)
ws.merge_cells(f'A{SUM_HDR}:D{SUM_HDR}'); ws.merge_cells(f'E{SUM_HDR}:J{SUM_HDR}')
ws[f'A{SUM_HDR}'] = '구 분'; ws[f'E{SUM_HDR}'] = '내 역'; ws[f'K{SUM_HDR}'] = '금액'
items = [
    ('1. 도입인프라(H/W)', 'AI 추론·RAG 문서구조화·AI-Ops 서버, 네트워크 — 견적 수령 후 별도 제출(본 견적가 미포함)', '별도'),
    (f'=A{SW_SEC}', 'DRM 보안SW 1식, SSO 인증서버SW 1식, STT 엔진 1식', f'=K{SW_TOT}'),
]
for (sec, hdr, first, tot) in dev_layout:
    items.append((f'=A{sec}', f'="IT PM 포함 "&TEXT(J{tot},"0.0")&"M/M"', f'=K{tot}'))
items.append(('단수 조정', '목표 견적가 기준 원 단위 조정', None))
for row, (a, e, k) in zip(SUM_ROWS, items):
    style_row(row, 'sum_hw' if row == SUM_ROWS[0] else 'sum_row', 34.2)
    ws.merge_cells(f'A{row}:D{row}'); ws.merge_cells(f'E{row}:J{row}')
    ws[f'A{row}'] = a; ws[f'E{row}'] = e; ws[f'K{row}'] = k
ws[f'K{SUM_ROWS[1]}'].alignment = copy.copy(ws[f'K{SUM_ROWS[2]}'].alignment)
from openpyxl.styles import Alignment
ws[f'K{SUM_ROWS[0]}'].alignment = Alignment(horizontal='right', vertical='center')
ADJ = SUM_ROWS[-1]
ws.row_dimensions[ADJ].height = 23.1
ws[f'K{ADJ}'] = f'=-MOD(SUM(K{SUM_ROWS[1]}:K{SUM_ROWS[-2]}),1000000)'
style_row(SUM_TOT, 'sum_tot', 23.1)
ws.merge_cells(f'A{SUM_TOT}:J{SUM_TOT}')
ws[f'A{SUM_TOT}'] = '총 계(부가가치세 별도)'
ws[f'K{SUM_TOT}'] = f'=SUM(K{SUM_ROWS[0]}:K{SUM_ROWS[-1]})'
ws['A11'] = f'="총 견적가 : 일금 "&NUMBERSTRING(K{SUM_TOT},1)&"원정 (\\"&TEXT(K{SUM_TOT},"#,###")&"원/VAT별도)"'

# ---- 1. HW (양식 유지, 별도 견적) ----
ws[f'A{HW_SEC}'] = '1. 도입인프라(H/W) - 견적 수령 후 별도 제출 (본 견적가 미포함)'
ws[f'A{HW_SEC}']._style = copy.copy(P['sec']['A'])
def table_hdr(row, labels):
    style_row(row, 'hw_hdr', 18.0)
    for m in ('A{0}:B{0}', 'C{0}:D{0}', 'E{0}:H{0}'):
        ws.merge_cells(m.format(row))
    for col, v in zip(('A', 'C', 'E', 'I', 'J', 'K'), labels):
        ws[f'{col}{row}'] = v
table_hdr(HW_HDR, ['용도', '제품명', '상세 내역', '납품단가', '수량', '제안가'])
for i, vals in enumerate(HW_ROWS):
    row = HW_FIRST + i
    style_row(row, 'hw_row', 49.2)
    for m in ('A{0}:B{0}', 'C{0}:D{0}', 'E{0}:H{0}'):
        ws.merge_cells(m.format(row))
    ws[f'A{row}'], ws[f'C{row}'], ws[f'E{row}'] = vals[0], vals[2], vals[4]
    ws[f'J{row}'] = vals[9]; ws[f'K{row}'] = f'=I{row}*J{row}'
style_row(HW_TOT, 'hw_tot', 23.1)
ws.merge_cells(f'A{HW_TOT}:J{HW_TOT}')
ws[f'A{HW_TOT}'] = '합계'; ws[f'K{HW_TOT}'] = f'=SUM(K{HW_FIRST}:K{HW_TOT-1})'

# ---- 2. 상용 SW ----
ws[f'A{SW_SEC}'] = '2. 상용 S/W 도입'
ws[f'A{SW_SEC}']._style = copy.copy(P['sec']['A'])
table_hdr(SW_HDR, ['용도', '제품명', '상세 내역', '납품단가', '수량', '제안가'])
for i, (use, prod, det, price, qty) in enumerate(SW):
    row = SW_FIRST + i
    style_row(row, 'hw_row', 36.0)
    for m in ('A{0}:B{0}', 'C{0}:D{0}', 'E{0}:H{0}'):
        ws.merge_cells(m.format(row))
    ws[f'A{row}'], ws[f'C{row}'], ws[f'E{row}'] = use, prod, det
    ws[f'I{row}'], ws[f'J{row}'] = price, qty
    ws[f'K{row}'] = f'=I{row}*J{row}'
    ws[f'K{row}'].number_format = ws[f'I{row}'].number_format
style_row(SW_TOT, 'hw_tot', 23.1)
ws.merge_cells(f'A{SW_TOT}:J{SW_TOT}')
ws[f'A{SW_TOT}'] = '합계'; ws[f'K{SW_TOT}'] = f'=SUM(K{SW_FIRST}:K{SW_TOT-1})'

# ---- 3~6. 개발비 ----
for n, ((title, label, roles), (sec, hdr, first, tot)) in enumerate(zip(DEV, dev_layout), 3):
    ws[f'A{sec}'] = f'{n}. 개발비({title})'
    ws[f'A{sec}']._style = copy.copy(P['sec']['A'])
    style_row(hdr, 'dv_hdr', 20.4)
    ws.merge_cells(f'A{hdr}:D{hdr}'); ws.merge_cells(f'E{hdr}:F{hdr}')
    for col, v in zip('AEGHIJK', ['구 분', '업무', '월평균 임금', '제경비(144%)', '기술료(20%)', '수량(M/M)', '제안가']):
        ws[f'{col}{hdr}'] = v
    for i, (role, task, mm) in enumerate(roles):
        row = first + i
        style_row(row, 'dv_first' if i == 0 else 'dv_row', 20.4)
        ws.merge_cells(f'C{row}:D{row}'); ws.merge_cells(f'E{row}:F{row}')
        ws[f'C{row}'], ws[f'E{row}'] = role, task
        ws[f'G{row}'] = RATE[role]
        ws[f'H{row}'] = f'=G{row}*144%'
        ws[f'I{row}'] = f'=(G{row}+H{row})*20%'
        ws[f'J{row}'] = mm
        ws[f'J{row}'].number_format = '0.0'
        ws[f'K{row}'] = f'=(G{row}+H{row}+I{row})*J{row}'
        ws[f'E{row}'].alignment = Alignment(horizontal='left', vertical='center', wrap_text=True, shrink_to_fit=False)
        if len(task) > 34:
            ws.row_dimensions[row].height = 30.0
    ws.merge_cells(f'A{first}:B{tot-1}')
    ws[f'A{first}'] = label
    style_row(tot, 'dv_tot', 22.2)
    ws.merge_cells(f'A{tot}:I{tot}')
    ws[f'A{tot}'] = '합계'
    ws[f'J{tot}'] = f'=SUM(J{first}:J{tot-1})'; ws[f'J{tot}'].number_format = '0.0'
    ws[f'K{tot}'] = f'=SUM(K{first}:K{tot-1})'

# ---- 관련 세부 사항 ----
notes = [
    '※ 기술협상안에 따라 개발비용은 조정이 가능합니다.',
    '※ 구축완료 후 S/W 1년, H/W 2년 무상하자보수 적용되며, 무상하자보수 종료 후 유상유지보수 전환 합니다',
    ' - S/W : S/W 도입가의 15%',
    ' - H/W : H/W 도입가의 12%',
    '※ 도입인프라(H/W)는 제조사 견적 수령 후 별도 제출하며, 본 총 견적가(10억원)에 포함되지 않습니다.',
    '※ 과제 3(대국민 서비스)의 외부망(DMZ)·웹방화벽 등 인프라 비용은 본 견적에 포함되지 않습니다.',
]
ws[f'A{NOTE}'] = '관련 세부 사항'; ws[f'A{NOTE}']._style = copy.copy(P['note_t']['A'])
ws.row_dimensions[NOTE].height = 14.4
for i, t in enumerate(notes, 1):
    ws[f'A{NOTE+i}'] = t; ws[f'A{NOTE+i}']._style = copy.copy(P['note']['A'])
    ws.row_dimensions[NOTE + i].height = 14.4
LAST = NOTE + len(notes) + 1
ws.print_area = f'A1:K{LAST}'
wb.calculation.fullCalcOnLoad = True
wb.save(OUT)
print('saved', OUT, 'rows', LAST, 'total row', SUM_TOT)
